from __future__ import annotations

import argparse
import csv
from pathlib import Path

from training_v2.common import extract_initials, normalize_text


def read_rows(path: Path) -> list[dict]:
    if path.suffix.lower() == ".xlsx":
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise SystemExit("XLSX를 읽으려면: pip install openpyxl") from exc
        wb = load_workbook(path, read_only=True, data_only=True)
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            return []
        headers = [str(x or "").strip() for x in rows[0]]
        out = []
        for values in rows[1:]:
            out.append({headers[i]: values[i] for i in range(min(len(headers), len(values)))})
        return out

    with path.open("r", encoding="utf-8-sig", newline="") as f:
        return list(csv.DictReader(f))


def main() -> None:
    p = argparse.ArgumentParser()
    p.add_argument("--input", required=True)
    p.add_argument("--output", default="training_v2/reviews/legacy_candidate_review.csv")
    args = p.parse_args()

    src = Path(args.input)
    rows = read_rows(src)
    unique: dict[tuple[str, str], dict] = {}

    for row in rows:
        initials = str(row.get("initials") or "").strip()
        split = str(row.get("split") or "dev").strip()
        candidates_raw = str(row.get("candidates") or "")
        approved = str(row.get("approved") or "1").strip()

        for text in [x.strip() for x in candidates_raw.split("|") if x.strip()]:
            got = extract_initials(text)
            if not initials:
                initials = got
            key = (initials, normalize_text(text))
            if key in unique:
                continue
            unique[key] = {
                "split": split,
                "initials": initials,
                "text": text,
                "extracted_initials": got,
                "source": f"legacy:{src.name}",
                # 기존 파일은 자동 학습에 넣지 않고 재검토하도록 기본 0.
                "approved": 0,
                "note": "" if got == initials else f"초성불일치:{got}",
            }

    dst = Path(args.output)
    dst.parent.mkdir(parents=True, exist_ok=True)
    with dst.open("w", encoding="utf-8-sig", newline="") as f:
        fields = ["split","initials","text","extracted_initials","source","approved","note"]
        w = csv.DictWriter(f, fieldnames=fields)
        w.writeheader()
        w.writerows(unique.values())

    print(f"MIGRATED UNIQUE CANDIDATES: {len(unique)}")
    print("OUTPUT:", dst)
    print("주의: legacy candidate는 approved=0 기본값입니다. 자연스러운 표현만 직접 1로 바꾸세요.")


if __name__ == "__main__":
    main()
