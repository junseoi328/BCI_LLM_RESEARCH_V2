from __future__ import annotations

import csv
from pathlib import Path


def load_review(path: str | Path) -> list[dict]:
    path = Path(path)
    if path.suffix.lower() == ".xlsx":
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise RuntimeError("XLSX review를 사용하려면 pip install openpyxl") from exc
        wb = load_workbook(path, read_only=True, data_only=True)
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            return []
        headers = [str(x or "").strip() for x in rows[0]]
        out = []
        for values in rows[1:]:
            out.append({headers[i]: values[i] for i in range(min(len(headers), len(values)))})
        return out

    with path.open("r", encoding="utf-8-sig", newline="") as f:
        return list(csv.DictReader(f))


def approved(value) -> bool:
    return str(value or "").strip().lower() in {"1","true","yes","y","approved","ok"}
